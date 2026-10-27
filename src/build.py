import json, os, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
os.makedirs(OUT, exist_ok=True)

SETTINGS = {"name": "Monster Tycoon", "start": 1500, "go": 200, "minutes": 60}
TEAMS = ["Team 1", "Team 2", "Team 3", "Team 4", "Team 5", "Team 6"]
NAMES = {
    "U1": "Science", "1A": "The Solar System", "1B": "The Human Body",
    "U2": "Geography", "2A": "UK Cities", "2B": "Rivers and Oceans", "2C": "World Capitals",
    "U3": "History", "3A": "The Tudors", "3B": "The World Wars", "3C": "Industrial Revolution",
    "U4": "Digital Skills", "4A": "Online Safety", "4B": "Computer Hardware", "4C": "The Internet",
    "U5": "Healthy Living", "5A": "Food and Nutrition", "5B": "First Aid", "5C": "Fitness",
    "U6": "Money Matters", "6A": "Budgeting", "6B": "Payslips", "6C": "Saving and Banking",
    "U7": "Environment", "7A": "Recycling", "7B": "Energy", "7C": "Climate",
    "U8": "Careers", "8A": "CVs", "8B": "Interviews",
    "HUB1": "Teamwork", "HUB2": "Communication", "HUB3": "Problem Solving", "HUB4": "Customer Service",
    "MATHS": "Maths", "ENGLISH": "English",
}
NAME_NOTES = {
    "U": "Unit (colour set) name", "HUB": "Skills hub (like a station)", "MATHS": "Utility square", "ENGLISH": "Utility square",
}

TF = ("True", "False", "", "")
# topic, level, question, A, B, C, D, answer, explanation
Q = [
 ("1A",1,"Which planet is closest to the Sun?","Venus","Mercury","Mars","Earth","B","Mercury has the smallest orbit of all the planets."),
 ("1B",1,"Which organ pumps blood around the body?","Lungs","Liver","Heart","Kidneys","C","The heart pumps blood through the arteries and veins."),
 ("2A",1,"What is the capital city of Wales?","Swansea","Cardiff","Newport","Bangor","B","Cardiff has been the capital of Wales since 1955."),
 ("2B",1,"What is the longest river in the UK?","Thames","Severn","Trent","Tyne","B","The Severn runs about 220 miles from Wales into the Bristol Channel."),
 ("2C",1,"What is the capital of Australia?","Sydney","Melbourne","Canberra","Perth","C","Canberra was purpose-built as the capital."),
 ("3A",1,"How many wives did Henry VIII have?","4","5","6","8","C","Henry VIII married six times."),
 ("3B",1,"In which year did the Second World War end?","1918","1939","1945","1950","C","The war ended in 1945."),
 ("3C",1,"Which invention powered many factories in the Industrial Revolution?","Steam engine","Electric motor","Petrol engine","Solar panel","A","Steam engines drove machines in mills and factories."),
 ("4A",1,"Which of these makes the strongest password?","Your pet's name","A long mix of words, numbers and symbols","Your birthday","password123","B","Long, random combinations are much harder to guess or crack."),
 ("4B",1,"What does CPU stand for?","Central Processing Unit","Computer Power Unit","Central Program Utility","Core Processing Upload","A","The CPU carries out a computer's instructions."),
 ("4C",1,"What does 'www' stand for in a web address?","World Wide Web","Wide World Web","Web World Wide","World Web Window","A","The World Wide Web was invented by Tim Berners-Lee."),
 ("5A",1,"How many portions of fruit and veg a day does the NHS recommend?","At least 3","At least 5","At least 7","At least 10","B","The NHS '5 A Day' guidance."),
 ("5B",1,"Which number do you call in a UK emergency?","101","111","999","911","C","999 (or 112) for emergencies; 111 is non-urgent medical help, 101 is non-emergency police."),
 ("5C",1,"How much moderate activity should adults aim for each week (NHS)?","30 minutes","75 minutes","150 minutes","300 minutes","C","At least 150 minutes of moderate activity, or 75 minutes vigorous."),
 ("6A",1,"What is a budget?","A plan for your income and spending","A type of bank account","A loan","A tax","A","A budget helps you see where your money goes."),
 ("6B",1,"On a UK payslip, what does 'NI' stand for?","National Income","National Insurance","Net Income","New Investment","B","National Insurance contributions are taken from your pay."),
 ("6C",1,"What is 'interest' on a savings account?","A fee you pay the bank","Money the bank pays you for saving","A type of tax","Your account number","B","Banks pay interest as a reward for keeping your money with them."),
 ("7A",1,"Which of these can usually go in household recycling?","Clean glass bottles","Nappies","Crisp packets","Used tissues","A","Clean glass bottles and jars are widely recycled."),
 ("7B",1,"Which of these is a renewable energy source?","Coal","Wind","Natural gas","Oil","B","Wind won't run out, unlike fossil fuels."),
 ("7C",1,"Which gas contributes most to human-caused climate change?","Oxygen","Nitrogen","Carbon dioxide","Helium","C","Burning fossil fuels releases carbon dioxide."),
 ("8A",1,"What should a CV include?","Your bank details","Your skills and experience","Your passwords","Your favourite film","B","A CV shows employers what you can offer."),
 ("8B",1,"What's a good thing to do at the end of an interview?","Ask a question about the role","Check your phone","Ask when you can book a holiday","Say nothing","A","Asking a question shows genuine interest."),
 ("U1",2,"Which gas do plants take in for photosynthesis?","Oxygen","Carbon dioxide","Nitrogen","Hydrogen","B","Plants use carbon dioxide, water and light to make glucose."),
 ("U1",2,"Which blood cells help fight infection?","Red blood cells","White blood cells","Platelets","Plasma","B","White blood cells are part of the immune system."),
 ("U1",3,"Why does Earth have seasons?","Its distance from the Sun changes","Its axis is tilted","The Moon's gravity","The Sun gets hotter in summer","B","The tilt changes how directly sunlight hits each half of the planet."),
 ("U2",2,"Which is the largest ocean?","Atlantic","Indian","Pacific","Arctic","C","The Pacific covers about a third of Earth's surface."),
 ("U2",2,"What is the highest mountain in the UK?","Yr Wyddfa (Snowdon)","Ben Nevis","Scafell Pike","Slieve Donard","B","Ben Nevis in Scotland is 1,345 m high."),
 ("U2",3,"Why were many cities built next to rivers?","Rivers are quieter","For water, trade and transport","Rivers stop flooding","It was the law","B","Rivers supplied water and were the main trade routes."),
 ("U3",2,"Who was the last Tudor monarch?","Henry VIII","Edward VI","Mary I","Elizabeth I","D","Elizabeth I died in 1603, ending the Tudor line."),
 ("U3",2,"In which year did the First World War begin?","1905","1914","1918","1939","B","The war began in 1914."),
 ("U3",3,"Why did many people move to towns during the Industrial Revolution?","To find work in factories","Farming was banned","Towns were cleaner","To avoid tax","A","New factories created jobs while farm work declined."),
 ("U4",2,"What is phishing?","A fake message trying to trick you into giving information","A type of computer fan","A fast internet connection","A way to back up files","A","Phishing messages pretend to be trusted organisations."),
 ("U4",2,"Which of these is an input device?","Monitor","Printer","Keyboard","Speaker","C","Input devices send data into the computer."),
 ("U4",3,"Why should you install software updates?","They make the screen brighter","They fix security weaknesses","They delete old files","They slow down hackers' internet","B","Updates patch known vulnerabilities."),
 ("U5",2,"What does CPR stand for?","Cardiopulmonary resuscitation","Central pulse recovery","Chest pressure routine","Cardiac pressure relief","A","CPR keeps blood and oxygen moving until help arrives."),
 ("U5",2,"Which nutrient is the body's main source of energy?","Carbohydrates","Vitamins","Minerals","Fibre","A","Carbohydrates are broken down into glucose for energy."),
 ("U5",3,"Why is warming up before exercise important?","It burns more calories","It prepares muscles and reduces injury risk","It makes you taller","It isn't important","B","Warm muscles are more flexible and less likely to strain."),
 ("U6",2,"What is the difference between gross and net pay?","Gross is after deductions","Net is after deductions","They're the same","Net includes bonuses","B","Net pay is what actually reaches your bank account."),
 ("U6",2,"What does APR show you?","The yearly cost of borrowing","Your account number","Your credit score","Your monthly salary","A","APR helps you compare loans and credit cards."),
 ("U6",3,"Why is an emergency fund useful?","It covers unexpected costs without borrowing","It earns the highest interest","It improves your payslip","It replaces insurance","A","Savings for emergencies stop surprise bills becoming debt."),
 ("U7",2,"What does 'carbon footprint' mean?","Total greenhouse gases caused by a person or activity","The size of your shoes","The amount of coal in the UK","A type of tree","A","It measures the climate impact of what we do."),
 ("U7",2,"Which light bulb uses the least energy?","Incandescent","Halogen","LED","Candle","C","LEDs use a fraction of the energy of older bulbs."),
 ("U7",3,"Why does reducing food waste help the environment?","It saves the energy, water and emissions used to produce food","Food waste is radioactive","It makes bins lighter","It doesn't help","A","Wasted food wastes everything that went into growing and transporting it."),
 ("U8",2,"What is a cover letter for?","Explaining why you're right for a specific job","Listing your hobbies only","Asking for a pay rise","Replacing your CV","A","A cover letter tailors your application to the role."),
 ("U8",2,"What does the STAR method help you do in an interview?","Answer with Situation, Task, Action, Result","Rate the interviewer","Remember names","Choose an outfit","A","STAR gives structure to examples of your skills."),
 ("U8",3,"Why should you research a company before an interview?","To show interest and tailor your answers","To find the manager's address","It isn't necessary","To ask for a higher salary","A","Research helps you link your skills to what they need."),
 ("HUB1",1,"A team member is struggling to finish their part. What's the best response?","Ignore it","Offer help or flag it early","Redo their work without telling them","Complain to others","B","Good teams support each other and raise problems early."),
 ("HUB2",1,"Which is an example of active listening?","Interrupting with your own story","Nodding and summarising what they said","Looking at your phone","Planning your reply while they talk","B","Summarising shows you've understood."),
 ("HUB3",1,"What's the first step in solving a problem?","Pick the quickest fix","Understand the problem clearly","Find someone to blame","Give up","B","You can't fix what you don't understand."),
 ("HUB4",1,"A customer is angry about a late order. What should you do first?","Argue back","Listen and acknowledge their frustration","Hang up","Blame the courier","B","Listening calms the situation before you solve it."),
 ("MATHS",1,"What is 15% of £80?","£8","£12","£15","£10","B","10% is £8, 5% is £4, so 15% is £12."),
 ("MATHS",1,"A train leaves at 14:45. The journey takes 1 hour 50 minutes. When does it arrive?","16:25","16:35","16:45","15:35","B","14:45 + 1 h = 15:45, + 50 min = 16:35."),
 ("MATHS",1,"What is 3/4 as a percentage?","34%","60%","75%","80%","C","3 ÷ 4 = 0.75 = 75%."),
 ("MATHS",1,"A room is 6 m by 4 m. What is its area?","10 m²","20 m²","24 m²","48 m²","C","Area = length × width = 24 m²."),
 ("ENGLISH",1,"Which word is spelled correctly?","Recieve","Receive","Receeve","Riceive","B","'I before e except after c'."),
 ("ENGLISH",1,"Fill the gap: '___ going to be late.'","Their","There","They're","Thier","C","They're = they are."),
 ("ENGLISH",1,"Which is the best greeting for a formal email to someone you don't know?","Hiya","Dear Sir or Madam","Yo","Alright?","B","Formal emails need a formal greeting."),
 ("ENGLISH",1,"Which punctuation mark ends a question?","Full stop","Exclamation mark","Question mark","Comma","C","Questions end with a question mark."),
 ("POPQUIZ",1,"True or false: the Great Wall of China is easily visible from the Moon.",*TF,"B","It's far too narrow to see from the Moon."),
 ("POPQUIZ",1,"True or false: lightning can strike the same place twice.",*TF,"A","Tall buildings are struck many times a year."),
 ("POPQUIZ",1,"True or false: bats are blind.",*TF,"B","Bats can see; many also use echolocation."),
 ("POPQUIZ",1,"True or false: botanically, a tomato is a fruit.",*TF,"A","It develops from the flower and contains seeds."),
 ("POPQUIZ",1,"True or false: goldfish only have a three-second memory.",*TF,"B","Studies show goldfish can remember things for months."),
 ("REALWORLD",1,"You're running 10 minutes late for work. What should you do?","Say nothing","Let your manager know as soon as possible","Go home","Post about it online","B","Telling your manager early shows reliability."),
 ("REALWORLD",1,"A colleague offers you their login to 'save time'. What should you do?","Use it","Politely refuse and use your own login","Share it with others","Write it on a sticky note","B","Sharing logins breaks security policies."),
 ("REALWORLD",1,"You make a mistake that affects a customer. What's the best approach?","Hide it","Own it, apologise and fix it","Blame a colleague","Wait to see if they notice","B","Owning mistakes builds trust."),
 ("REALWORLD",1,"An email from 'IT' urgently asks for your password. What do you do?","Reply with it","Report it as phishing","Forward it to friends","Ignore all IT emails forever","B","Real IT teams never ask for your password."),
 ("REALWORLD",1,"Your manager gives feedback you disagree with. Best response?","Walk out","Listen, ask questions and discuss calmly","Ignore it","Complain online","B","Calm discussion helps you understand and improve."),
 ("REVISION",1,"What is 7 × 8?","54","56","64","48","B","7 × 8 = 56."),
 ("REVISION",1,"What is the capital of Scotland?","Glasgow","Edinburgh","Aberdeen","Dundee","B","Edinburgh is Scotland's capital."),
 ("REVISION",1,"How many minutes are in 2.5 hours?","120","130","150","250","C","2 × 60 + 30 = 150."),
 ("REVISION",1,"What is the plural of 'child'?","Childs","Children","Childes","Childrens","B","'Children' is an irregular plural."),
 ("BONUS",1,"What is the largest planet in the Solar System?","Saturn","Jupiter","Neptune","Earth","B","Jupiter is more than twice as massive as all the other planets combined."),
 ("BONUS",1,"How many sides does a hexagon have?","5","6","7","8","B","Hex means six."),
 ("BONUS",1,"Which city hosted the 2012 Olympic Games?","Beijing","London","Rio de Janeiro","Athens","B","London hosted the 2012 Summer Olympics."),
 ("BONUS",1,"What is H2O more commonly called?","Water","Salt","Hydrogen peroxide","Oxygen","A","Two hydrogen atoms and one oxygen atom make water."),
]

questions = []
for n, (topic, lvl, q, a, b, c, d, ans, exp) in enumerate(Q):
    opts = [{"L": L, "text": t} for L, t in zip("ABCD", (a, b, c, d)) if t]
    questions.append({"id": f"s{n}", "topic": topic, "level": lvl, "q": q, "opts": opts, "ans": ans, "exp": exp})

data = {"settings": SETTINGS, "teams": TEAMS, "names": NAMES, "questions": questions}

# ---- HTML ----
html = open(os.path.join(os.path.dirname(os.path.abspath(__file__)),"index.template.html"), encoding="utf-8").read()
assert "/*__SAMPLE_DATA__*/null" in html
html = html.replace("/*__SAMPLE_DATA__*/null", json.dumps(data, ensure_ascii=False))
open(f"{OUT}/index.html", "w", encoding="utf-8").write(html)

# ---- Excel ----
wb = Workbook()
H = Font(bold=True, color="FFFFFF", size=12)
HF = PatternFill("solid", fgColor="22305A")
NOTE = Font(italic=True, color="555555")
def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.cell(1, i).font = H; ws.cell(1, i).fill = HF
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"

ws = wb.active; ws.title = "How to use"
ws.column_dimensions["A"].width = 110
lines = [
 ("How to use this question bank", Font(bold=True, size=16)),
 ("1. Settings: game name, starting money, money for passing GO, default minutes.", None),
 ("2. Names: rename the 8 units (U1–U8), the 22 property topics (1A–8B), the 4 skills hubs and the Maths/English squares.", None),
 ("3. Questions: one row per question. Topic must be a code from the Names sheet, or one of: HUB, POPQUIZ, REALWORLD, REVISION, BONUS, ANY.", None),
 ("   Level 1 = buying a property. Level 2 = houses 1–2. Level 3 = houses 3–4. (Hotels are a 'teach it back' challenge, no question needed.)", None),
 ("   Put level 2–3 questions on the unit code (e.g. U3) so they're shared by every topic in that unit, or on a topic code (e.g. 3A).", None),
 ("   Answer: the letter A–D. For open questions, leave the options blank and type the answer text instead.", None),
 ("   Pop quiz: use True/False in options A and B.", None),
 ("4. Missing questions? The game falls back: topic → unit → ANY. Aim for at least 2 questions per topic.", None),
 ("5. Save the file, then in the game press 'Load Excel file' on the setup screen.", None),
 ("   On GitHub Pages, a file named questions.xlsx next to index.html loads automatically.", None),
 ("", None),
 ("Don't rename the sheets or the column headings — the game looks for them by name.", Font(bold=True, color="A3171B")),
]
for t, f in lines:
    ws.append([t])
    if f: ws.cell(ws.max_row, 1).font = f

ws = wb.create_sheet("Settings"); header(ws, ["Setting", "Value"], [22, 40])
for k, v in [("Game name", SETTINGS["name"]), ("Starting money", 1500), ("GO money", 200), ("Minutes", 60)]:
    ws.append([k, v])


ws = wb.create_sheet("Names"); header(ws, ["Code", "Name", "What it is"], [10, 34, 40])
for code, name in NAMES.items():
    if code.startswith("U"): note = "Unit (colour set) name"
    elif code.startswith("HUB"): note = "Skills hub (like a station)"
    elif code in ("MATHS", "ENGLISH"): note = "Functional skills square (like a utility)"
    else: note = f"Property in Unit {code[0]}"
    ws.append([code, name, note])
    if code.startswith("U"):
        for c in range(1, 4): ws.cell(ws.max_row, c).font = Font(bold=True)

ws = wb.create_sheet("Questions")
header(ws, ["Topic", "Level", "Question", "Option A", "Option B", "Option C", "Option D", "Answer", "Explanation"], [10, 7, 60, 28, 28, 28, 28, 8, 50])
for q in Q:
    ws.append(list(q))
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True); ws.add_data_validation(dv); dv.add("B2:B1000")

wb.save(f"{OUT}/questions.xlsx")
print("questions:", len(questions))
